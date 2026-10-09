#!/usr/bin/env python3
"""Render docs/DEVLOG.md as a Word document.

The markdown file stays the source of truth — it diffs cleanly in git and sits
next to the code it describes. This produces the .docx for reading, printing and
handing to a supervisor, and is safe to re-run: the output is overwritten.

Run with `make devlog`.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parent.parent
SOURCE = REPO_ROOT / "docs" / "DEVLOG.md"
OUTPUT = REPO_ROOT / "docs" / "Sendox-Development-Log.docx"

MONO = "Consolas"
BODY = "Calibri"
CODE_GREY = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x60, 0x60, 0x60)

# `code`, **bold**, *italic* — parsed in one pass so nesting cannot confuse order.
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_runs(paragraph, text: str) -> None:
    """Write text into a paragraph, honouring inline markdown."""
    # Links become just their label; a .docx full of raw URLs reads badly.
    text = LINK.sub(r"\1", text)

    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
            run.font.size = Pt(9.5)
            run.font.color.rgb = CODE_GREY
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def add_code_block(document: Document, lines: list[str], language: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)

    if language:
        label = paragraph.add_run(f"{language}\n")
        label.font.name = MONO
        label.font.size = Pt(8)
        label.font.color.rgb = MUTED

    run = paragraph.add_run("\n".join(lines))
    run.font.name = MONO
    run.font.size = Pt(9)
    run.font.color.rgb = CODE_GREY

    # Shade the block so it reads as code rather than indented prose.
    shading = paragraph._p.get_or_add_pPr().makeelement(
        qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "F4F4F6"}
    )
    paragraph._p.get_or_add_pPr().append(shading)


def add_table(document: Document, rows: list[list[str]]) -> None:
    header, *body = rows
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for cell, text in zip(table.rows[0].cells, header, strict=False):
        cell.paragraphs[0].text = ""
        add_runs(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for raw in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, raw, strict=False):
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], text)

    document.add_paragraph()


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def convert(markdown: str) -> Document:
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)

    lines = markdown.splitlines()
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            language = stripped[3:].strip()
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code_block(document, block, language)
            index += 1
            continue

        # tables: a header row followed by a |---| separator
        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and set(lines[index + 1].strip()) <= set("|-: ")
            and "-" in lines[index + 1]
        ):
            rows = [split_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        if not stripped:
            index += 1
            continue

        if stripped == "---":
            rule = document.add_paragraph()
            rule.paragraph_format.space_before = Pt(10)
            run = rule.add_run("· · ·")
            run.font.color.rgb = MUTED
            rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = document.add_heading(level=min(level, 4))
            heading.text = ""
            add_runs(heading, stripped[level:].strip())
            index += 1
            continue

        # numbered list
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, numbered.group(2))
            index += 1
            continue

        # bullet, including a continuation line indented beneath it
        if stripped.startswith(("- ", "* ")):
            text = stripped[2:]
            index += 1
            while (
                index < len(lines)
                and lines[index].startswith("  ")
                and lines[index].strip()
                and not lines[index].strip().startswith(("-", "*", "|", "#", "```"))
            ):
                text += " " + lines[index].strip()
                index += 1
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, text)
            continue

        # paragraph: join the wrapped lines back together
        text = stripped
        index += 1
        while index < len(lines) and lines[index].strip() and not lines[index].strip().startswith(
            ("#", "-", "*", "|", "```", "---")
        ) and not re.match(r"^\d+\.\s", lines[index].strip()):
            text += " " + lines[index].strip()
            index += 1
        paragraph = document.add_paragraph()
        add_runs(paragraph, text)

    return document


def main() -> int:
    if not SOURCE.exists():
        print(f"missing {SOURCE}", file=sys.stderr)
        return 1

    document = convert(SOURCE.read_text(encoding="utf-8"))
    document.save(OUTPUT)
    size_kb = OUTPUT.stat().st_size / 1024
    print(f"wrote {OUTPUT.relative_to(REPO_ROOT)} ({size_kb:.0f} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
